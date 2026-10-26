"""Builds report/Project7_BMS_IoT_Report.docx from the results/ and figures/ folders.

Run after run_demo.py and the analysis scripts:  python report/build_report.py
"""
import csv
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from config import FLAG_TEXT  # noqa: E402
RES, FIG = ROOT / "results", ROOT / "figures"
OUT = Path(__file__).with_name("Project7_BMS_IoT_Report.docx")


def rows(name):
    with open(RES / name) as f:
        return list(csv.DictReader(f))


summary = json.loads((RES / "summary.json").read_text())
lat = summary["latency_ms"]
gw = summary["gateway"]
prot = rows("protection_table.csv")
canmsg = rows("can_message_table.csv")
pm = {r["protocol"]: float(r["transfer_ms"]) for r in rows("protocol_metrics.csv")}
with open(RES / "mqtt_qos.csv") as f:
    qos_lines = [l for l in csv.reader(f)]
qos_lat = qos_lines[1:4]
qos_out = qos_lines[6:9]
load500 = sum(float(r["load_500k_%"]) for r in canmsg)
load125 = sum(float(r["load_125k_%"]) for r in canmsg)

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.54)
    s.left_margin = s.right_margin = Cm(2.54)

fig_no = {}
tab_no = {}


def para(text="", bold=False, align=None, size=None, italic=False, after=None, keep=False):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic
        if size:
            r.font.size = Pt(size)
    if align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "j" or align is None:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def rich(parts):
    """parts: list of (text, bold) -> one justified paragraph."""
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for t, b in parts:
        p.add_run(t).bold = b
    return p


def heading(text, level=1):
    h = doc.add_heading(level=level)
    r = h.add_run(text)
    r.font.name = "Times New Roman"; r.font.color.rgb = RGBColor(0, 0, 0)
    r.font.size = Pt(14 if level == 1 else 12); r.bold = True
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    h.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    h.paragraph_format.space_after = Pt(6)
    return h


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if isinstance(it, tuple):
            p.add_run(it[0]).bold = True; p.add_run(it[1])
        else:
            p.add_run(it)
        p.paragraph_format.space_after = Pt(3)


def equation(text, n):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{text}          ({n})"); r.italic = False
    r.font.name = "Cambria Math"
    return p


def shade(cell, color):
    tc = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), color); tc.append(sh)


def table(key, caption, header, data, widths, size=10, note=None):
    tab_no[key] = f"{len(tab_no) // 1 + 0}"
    sec = current_section[0]
    k = sum(1 for v in tab_no.values() if v.startswith(f"{sec}.")) + 1
    tab_no[key] = f"{sec}.{k}"
    para(f"Table {sec}.{k}: {caption}", bold=True, align="c", size=11, after=3, keep=True)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size)
        shade(c, "D9E2F3")
    for row in data:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
        tr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); tr.append(cs)
    if note:
        para(note, italic=True, size=9, after=8)
    else:
        para(after=2)
    return t


def figure(key, path, caption, width=15.5):
    sec = current_section[0]
    k = sum(1 for v in fig_no.values() if v.startswith(f"{sec}.")) + 1
    fig_no[key] = f"{sec}.{k}"
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width))
    para(f"Figure {sec}.{k}: {caption}", align="c", size=10, italic=True, after=10)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def code(text):
    for line in text.strip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line if line else " "); r.font.name = "Consolas"; r.font.size = Pt(8.5)


current_section = [0]


def section(n, title):
    current_section[0] = n
    heading(f"{n}. {title}")


# ---------------------------------------------------------------- page numbers
def add_page_number(sec):
    p = sec.header.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
        r._r.append(el)


add_page_number(doc.sections[0])

# ================================================================ COVER
for _ in range(2):
    para()
para("PROJECT 7: BMS COMMUNICATION AND IoT-BASED BATTERY MONITORING DASHBOARD",
     bold=True, align="c", size=16, after=30)
for name in ["Student Name 1 (Roll No.)", "Student Name 2 (Roll No.)", "Student Name 3 (Roll No.)",
             "Student Name 4 (Roll No.)", "Student Name 5 (Roll No.)"]:
    para(name, align="c", size=13, after=2)
para(after=24)
para("23RF16 – Battery Management Systems", bold=True, align="c", size=13, after=90)
para("October 2026", align="c", size=12, after=40)
para("DEPARTMENT OF ROBOTICS AND AUTOMATION ENGINEERING", bold=True, align="c", size=13, after=10)
para("PSG COLLEGE OF TECHNOLOGY", bold=True, align="c", size=13, after=10)
para("COIMBATORE - 641004", bold=True, align="c", size=13)
page_break()

# ================================================================ ABSTRACT
heading("ABSTRACT")
para(
    "A Battery Management System (BMS) is only useful if the information it measures reaches the "
    "vehicle controller, the charger and the people who look after the battery. This project "
    "simulates the complete communication chain of a BMS, from the cells to an Internet-of-Things "
    "(IoT) dashboard, in Python. A 12S3P lithium-ion pack (44.4 V, 7.5 Ah) with a lumped thermal "
    "model is driven by a drive-cycle-like load, and a simulated BMS ECU measures cell voltages, "
    "current and temperatures, detects faults and opens the contactor. The BMS sends eight CAN "
    "messages, defined in a DBC file, on a python-can virtual bus. A gateway decodes the frames, "
    "checks their rolling counters and checksums, detects loss of communication and publishes JSON "
    "telemetry and alarm events over MQTT to a web dashboard and a Node-RED dashboard.")
para(
    f"A 50-minute test scenario injected seven events: over-voltage during regenerative braking, "
    f"over-current, CAN communication loss, an open temperature sensor, a short circuit, low SoC "
    f"and an early thermal-runaway trend. Every event was detected and shown on the dashboard. "
    f"The BMS used {load500:.2f} % of a 500 kbit/s CAN bus, and the median delay from a CAN frame "
    f"to the dashboard was {lat['p50']:.1f} ms. MQTT QoS 1 delivered all alarms sent during a 3 s "
    f"network outage while QoS 0 lost them. CAN, SMBus, UART, Bluetooth LE and LoRa were compared "
    f"for BMS use: CAN suits the on-board protection network, while BLE and LoRa suit monitoring "
    f"only.")
page_break()

# ================================================================ 1 INTRODUCTION
section(1, "INTRODUCTION")
para(
    "A BMS measures every cell voltage, the pack current and several temperatures, estimates the "
    "state of charge (SoC) and state of health (SoH), and protects the pack by limiting current or "
    "opening the main contactor [1], [2]. None of this is useful unless it is communicated. Inside "
    "a vehicle the BMS reports to the vehicle control unit, the motor controller and the charger, "
    "almost always over the Controller Area Network (CAN) [3], [4]. Smaller products use SMBus "
    "(laptop batteries), UART or Bluetooth Low Energy (e-bikes, power tools).")
para(
    "In recent years battery data has also been sent off the vehicle. Fleet operators, energy-storage "
    "sites and battery-swapping stations use IoT telemetry to see the state of every pack remotely, "
    "to receive alarms and to plan maintenance. A common way to do this is a gateway that reads the "
    "CAN bus and publishes the data with the lightweight MQTT protocol [6] to a dashboard. This "
    "project builds and tests such a chain entirely in simulation.")

section(2, "PROBLEM STATEMENT")
para(
    "The project brief (Project 7, syllabus unit: BMS Communication and IoT) [15] requires CAN bus "
    "BMS messages (cell voltages, SoC, temperature and fault flags) to be simulated with python-can "
    "and a virtual bus. The battery telemetry must be published over MQTT to a Node-RED or "
    "ThingSpeak dashboard with alarms, and CAN, SMBus, UART, Bluetooth LE and LoRa must be "
    "compared for BMS use. The suggested team split is: battery data generator, CAN message "
    "design, MQTT pipeline, dashboard and protocol comparison. The project is simulation-only and "
    "no hardware is required.")

section(3, "OBJECTIVES")
bullets([
    "Build a battery data generator: a 12-cell pack model with cell-to-cell variation, a thermal "
    "model, a load profile and injected faults.",
    "Design the BMS CAN message set in a DBC file and transmit it with python-can on a virtual bus.",
    "Implement the BMS protection logic (warning and fault levels, contactor control) so that the "
    "fault flags on the bus are meaningful.",
    "Build a CAN-to-MQTT gateway with data-integrity checks, a communication-loss watchdog and "
    "alarm events.",
    "Build a live dashboard (web page and Node-RED) with gauges, trends and alarms.",
    "Measure detection time, end-to-end latency, CAN bus load and MQTT delivery, and compare "
    "CAN, SMBus, UART, BLE and LoRa for BMS use.",
])

section(4, "SCOPE OF THE PROJECT")
para(
    "The project is a software simulation at pack level. The CAN bus is the python-can virtual "
    "bus, so electrical effects (bus termination, noise, bit errors) are not reproduced by the "
    "bus itself; their effect on data integrity is studied separately with a bit-level CAN frame "
    "model. The MQTT broker, gateway and dashboards run on one computer, so network delays of a "
    "real cellular or Wi-Fi link are not included. The battery model is simple (one RC pair per "
    "cell group, lumped thermal nodes) because the focus is on communication, not on cell "
    "modelling. The same code can run against real hardware by changing the python-can interface "
    "(for example to a USB-CAN adapter) and the broker address.")

section(5, "SYSTEM REQUIREMENTS")
table("req", "System requirements and simulation parameters", ["Component / Parameter", "Requirement"], [
    ["Programming language", "Python 3.13 (JavaScript for the dashboard and broker)"],
    ["Python libraries", "NumPy, Matplotlib, python-can 4.6, cantools, paho-mqtt 2.1"],
    ["IoT tools", "Aedes MQTT broker (Node.js), Node-RED 4 with node-red-dashboard, web browser"],
    ["Hardware", "None (simulation only); any laptop or desktop"],
    ["Battery pack", "12S3P Li-ion NMC 18650 (2.5 Ah cells): 44.4 V nominal, 7.5 Ah, 333 Wh"],
    ["Cell model", "1-RC Thevenin per group: R0 = 13 mΩ, R1 = 8 mΩ, τ = 60 s; 15-point OCV curve"],
    ["Cell variation", "Capacity ±1.5 %, R0 ±5 %; cell 3 weak (90 % capacity, +25 % R0)"],
    ["Thermal model", "C = 135 J/K per group, 0.25 W/K to ambient (30 °C), 0.30 W/K to neighbours"],
    ["Sensors", "12 cell voltages (±2 mV, 1 mV resolution), current (σ = 50 mA), 4 NTC sensors"],
    ["CAN bus", "python-can virtual bus, 500 kbit/s assumed, 11-bit IDs, 8 messages"],
    ["MQTT", "Broker on localhost: TCP 1883 (1884 if busy), WebSocket 9001; telemetry 1 Hz"],
    ["Test scenario", "3000 s: drive cycle 90 % → 10 % SoC, rest, charge at 5 A; 7 injected events"],
], [5.0, 11.0])

section(6, "WORKING PRINCIPLE")
figure("arch", FIG / "architecture.png",
       "System architecture: battery model, BMS ECU, CAN bus, gateway, MQTT broker and dashboards")
para(
    f"Figure {fig_no['arch']} shows the system. Every 100 ms the battery model computes the true "
    "cell voltages and temperatures. The BMS ECU adds sensor noise, evaluates the protection "
    "limits and encodes the results into CAN frames. The gateway receives the frames, decodes them "
    "with the same DBC file, and publishes the data over MQTT. Any number of dashboards can "
    "subscribe to the broker; the BMS and the dashboards never talk to each other directly.")
heading("6.1 Battery data generator", 2)
para("Each series group (three cells in parallel) is modelled with SoC z, polarisation voltage v1 "
     "and temperature T. With current i positive on discharge and time step Δt = 0.1 s:")
equation("z(k+1) = z(k) − Δt · i(k) / (3600 · Q)", 1)
equation("v(k) = OCV(z) − v1(k) − R0 · i(k)", 2)
equation("C · dT/dt = i²R0 + v1·i + P_fault − h(T − T_amb) − Σ h_n(T − T_neighbour)", 3)
para("P_fault is extra heat injected to simulate an internal short circuit. Faults such as a sensor "
     "open circuit or a stop of CAN transmission are injected in the BMS and bus layer.")
heading("6.2 CAN bus", 2)
para(
    "CAN is a multi-master differential bus. A frame with a lower identifier wins arbitration, so "
    "IDs set the message priority. Each frame carries up to 8 data bytes and is protected by a "
    "15-bit CRC, bit stuffing and acknowledgement [3], [4]. A standard 8-byte frame is 111 bits "
    "plus up to 24 stuff bits. The bus load is the share of time the bus is busy:")
equation("Bus load = Σ (frame bits_j / cycle time_j) / bit rate × 100 %", 4)
heading("6.3 MQTT", 2)
para(
    "MQTT is a publish/subscribe protocol over TCP [6]. Clients publish messages to named topics "
    "on a broker, and the broker forwards them to every subscriber of the topic. Three quality-of-"
    "service levels exist: QoS 0 (at most once), QoS 1 (at least once, acknowledged) and QoS 2 "
    "(exactly once, four-packet handshake). A retained message is kept by the broker for new "
    "subscribers, and a 'last will' message is published by the broker if a client disconnects "
    "unexpectedly. The dashboard uses MQTT over WebSocket so that it can run in a browser.")

section(7, "METHODOLOGY")
table("meth", "Steps followed in the project", ["Step", "Description"], [
    ["Pack model", "12 cell groups with 1-RC model, cell-to-cell variation and lumped thermal model."],
    ["Load and faults", "Drive-cycle current with regeneration, a rest, a charge, and 7 scheduled events."],
    ["BMS logic", "Warning/fault thresholds with debounce, latching faults, contactor and state machine."],
    ["CAN design", "8 messages in a DBC file; rolling counter and checksum on safety messages."],
    ["CAN transmission", "Frames sent cyclically (and on change of fault flags) on a python-can virtual bus."],
    ["Gateway", "Decode with cantools, check integrity, watchdog, alarms, JSON over MQTT."],
    ["Dashboards", "Web dashboard (Chart.js, MQTT.js) and Node-RED dashboard with alarm table."],
    ["Recording", "A subscriber logs every MQTT message with its arrival time for analysis."],
    ["Analysis", "Detection time, latency, bus load, MQTT QoS test, protocol comparison."],
], [3.6, 12.4])
table("team", "Team split", ["Member", "Responsibility"], [
    ["Member 1", "Battery data generator: pack, thermal model, load profile, fault injection"],
    ["Member 2", "CAN message design: DBC file, BMS ECU encoding, bus-load and CRC analysis"],
    ["Member 3", "MQTT pipeline: gateway, integrity checks, alarms, broker, QoS tests"],
    ["Member 4", "Dashboards: web dashboard and Node-RED flow"],
    ["Member 5", "Protocol comparison (CAN, SMBus, UART, BLE, LoRa) and report"],
], [3.0, 13.0], note="Replace the member labels with the team members' names.")

section(8, "CAN MESSAGE DESIGN")
para(
    "The message set is defined in can_db/bms.dbc so that the BMS and the gateway use exactly the "
    "same signal layout. All signals are little-endian (Intel). Fast-changing, safety-relevant "
    "data uses the lowest IDs and the shortest cycle; cell voltages are sent every 200 ms and "
    "temperatures every second.")
table("canset", "BMS CAN message set (11-bit identifiers, 8 data bytes each)",
      ["ID", "Message", "Cycle", "Signals (scaling)"], [
    ["0x100", "BMS_PackStatus", "100 ms", "Pack voltage (0.01 V), current (0.1 A, signed), SoC (0.5 %), SoH, "
     "BMS state, rolling counter, checksum"],
    ["0x101", "BMS_CellSummary", "100 ms", "Min/max cell voltage (1 mV) and cell index, max/min temperature"],
    ["0x102", "BMS_Limits", "100 ms", "Discharge/charge current limits, pack power, contactor state, charger, "
     "balancing"],
    ["0x110–0x112", "BMS_CellV_xx_yy", "200 ms", "4 cell voltages per frame (1 mV) – 12 cells in 3 frames"],
    ["0x120", "BMS_Temperatures", "1000 ms", "4 sensor temperatures (0.1 °C, signed)"],
    ["0x130", "BMS_Faults", "100 ms + on change", "16-bit warning flags, 16-bit fault flags, balancing "
     "mask, rolling counter, fault count, checksum"],
], [2.3, 3.3, 2.3, 8.1], size=9.5)
para(
    "BMS_PackStatus and BMS_Faults end with a 4-bit rolling counter and an 8-bit checksum "
    "(0xFF XOR the sum of bytes 0–6 and the low byte of the ID). The receiver can then detect a "
    "frozen sender (counter not changing), lost frames (counter jumps) and corrupted or "
    "misrouted data, in the same way as automotive end-to-end protection [7]. BMS_Faults is also "
    "sent immediately when a flag changes, so an alarm does not wait for the next cycle.")
table("flags", "Protection limits and fault-flag bits",
      ["Bit", "Flag", "Warning", "Fault / action"], [
    ["0", "Cell over-voltage", "> 4.20 V", "> 4.25 V – open contactor"],
    ["1", "Cell under-voltage", "< 3.10 V", "< 2.80 V – open contactor"],
    ["2", "Discharge over-current", "> 25 A", "> 35 A for 1 s – open contactor"],
    ["3", "Charge over-current", "> 12 A", "> 18 A for 1 s – open contactor"],
    ["4", "Over-temperature", "> 50 °C", "> 60 °C – open contactor"],
    ["6", "Cell imbalance (at rest)", "spread > 80 mV", "–"],
    ["7", "Low SoC", "< 15 %", "–"],
    ["8", "Thermal-runaway trend", "dT/dt > 1.5 °C/min above 40 °C (latched)", "–"],
    ["9", "Short circuit", "–", "> 100 A, instantaneous, latched until reset"],
    ["10", "Temperature sensor fault", "reading outside −35…120 °C", "sensor ignored"],
], [1.1, 4.4, 5.2, 5.3], size=9.5,
      note="Warnings need 0.5 s and faults 1 s of persistence (debounce). Bit 5 is under-temperature (0 °C / −10 °C).")

section(9, "MQTT PIPELINE AND DASHBOARD")
heading("9.1 Gateway", 2)
para(
    "The gateway (gateway.py) is the link between the vehicle network and the cloud. For every "
    "frame it looks up the message in the DBC, decodes the signals and checks the counter and "
    "checksum. It keeps the latest value of every signal and publishes one JSON telemetry message "
    "per second. When a bit in the warning or fault field changes it publishes an alarm event "
    "immediately. A watchdog raises a 'BMS communication lost' alarm if no BMS_PackStatus frame "
    "has arrived for 0.5 s. Optionally, the gateway also uploads eight fields to a ThingSpeak "
    "channel every 15 s (the free-tier limit) when an API key is set.")
table("topics", "MQTT topics", ["Topic", "QoS", "Content"], [
    ["bms/pack01/telemetry", "0", "JSON, 1 Hz: pack V, I, SoC, power, state, contactor, 12 cells, 4 temperatures, "
     "limits, flags"],
    ["bms/pack01/alarms", "1", "Event: code, level (WARNING/FAULT), raised or cleared, text, source, time"],
    ["bms/pack01/status", "1, retained", "Gateway online/offline; 'offline' is also the last-will message"],
    ["bms/pack01/gateway", "0", "Gateway statistics every 10 s: frames, checksum errors, counter jumps"],
], [4.6, 2.2, 9.2], size=9.5)
heading("9.2 Dashboards", 2)
para(
    "Two dashboards subscribe to the same topics. The web dashboard (dashboard/index.html) "
    "connects to the broker over WebSocket and shows SoC, pack voltage, current, power, maximum "
    "temperature, cell spread, contactor state and current limits, a voltage/current trend, a bar "
    "chart of the 12 cell voltages (cells outside limits in red), the temperature trend, an "
    "active-alarm list, an alarm log and link statistics. A red banner and a short beep appear "
    "when a fault is active. The Node-RED flow (nodered/flows.json) builds a similar dashboard "
    "with gauges, a chart, an alarm table and a pop-up for faults, as suggested in the brief.")

section(10, "TEST SCENARIOS AND PERFORMANCE METRICS")
table("tests", "Fault-injection test cases (simulation time)", ["Time (s)", "Event", "Expected BMS reaction"], [
    ["75–81", "Regenerative braking at 90 % SoC (−14 A)", "Over-voltage and charge over-current warnings"],
    ["600–607", "Load surge: 30 A for 4 s, then 40 A", "Over-current warning, then fault and contactor open"],
    ["1000–1005", "BMS stops transmitting on CAN for 5 s", "Gateway raises communication-loss alarm"],
    ["1300–1500", "NTC sensor 2 open circuit (reads −40 °C)", "Sensor-fault warning, sensor ignored"],
    ["1700", "External short circuit, 150 A pulse", "Instant fault, contactor latched open until reset"],
    ["1760", "Service reset", "Latched fault cleared, contactor closes again"],
    ["~2100", "Low SoC, weak cell 3 under load", "Low-SoC and under-voltage warnings"],
    ["2400–3000", "Module-3 fan failure and growing micro-short in cell 8 (2 → 16 W) while charging",
     "Thermal-runaway-trend warning, over-temperature warning, then fault"],
], [2.3, 7.0, 6.7], size=9.5)
para("The following metrics were measured:")
bullets([
    ("Detection delay: ", "time from the start of an event to the BMS flag on the bus."),
    ("End-to-end latency: ", "time from a CAN frame being sent to the MQTT message arriving at a "
     "subscriber (real time, measured on the same computer clock)."),
    ("CAN bus load: ", "Equation (4), with the exact number of stuff bits of each real frame."),
    ("MQTT delivery: ", "latency and number of messages lost for QoS 0, 1 and 2, with and without "
     "a 3 s subscriber outage."),
    ("Protocol comparison: ", "time to transfer one 48-byte BMS snapshot (12 cell voltages, 4 "
     "temperatures, status and flags) over each link, plus range, topology and typical use."),
])
para("The full scenario was run at ten times real speed (300 s) with all dashboards connected.")

section(11, "RESULTS AND OUTPUT")
heading("11.1 Telemetry and fault detection", 2)
figure("tel", FIG / "telemetry_overview.png",
       "Telemetry received over MQTT: (a) pack voltage and current, (b) SoC, (c) min/max cell "
       "voltage, (d) temperatures. Shaded bands mark the injected events.", width=15.0)
para(
    f"Figure {fig_no['tel']} shows the data exactly as the dashboard received it. The pack "
    "discharged from 90 % to about 10 % SoC during the drive cycle. The weak cell 3 was always the "
    "lowest and reached the under-voltage warning first near the end of the drive. The current "
    "drops to zero during the over-current and short-circuit faults because the BMS opened the "
    "contactor. During charging, sensor T3 (cells 7–9) rose steadily after the fan failure and "
    "micro-short, while the other sensors cooled down.")
prot_rows = []
for r in prot:
    prot_rows.append([r["event"].replace("_", " ").capitalize().replace("Can ", "CAN ").replace("soc", "SoC"),
                      f"{FLAG_TEXT[r['flag']]} ({r['level'].lower()})", r["injected_s"], r["bms_detect_s"],
                      r["detection_delay_s"], f"{float(r['dashboard_latency_ms']):.1f}",
                      r["cleared_s"]])
table("prot", "Protection response: injected event, detection and alarm delivery",
      ["Event", "Flag raised", "Injected (s)", "Detected (s)", "Delay (s)", "To dashboard (ms)",
       "Cleared (s)"], prot_rows, [3.0, 3.6, 1.7, 1.8, 1.4, 2.0, 2.5], size=8.5,
      note="Delay = detection time − event start. 'To dashboard' = CAN frame → MQTT alarm at the subscriber.")
tr = next(r for r in prot if r["flag"] == "TR_TREND")
otw = next(r for r in prot if r["flag"] == "OVER_TEMP" and r["level"] == "WARNING")
otf = next(r for r in prot if r["flag"] == "OVER_TEMP" and r["level"] == "FAULT")
para(
    f"All injected events were detected (Table {tab_no['prot']}). Threshold events were flagged "
    "0.4 s after they started, which is the 0.5 s warning debounce at 100 ms sampling. The "
    "over-current fault appeared 4.9 s after the start of the surge because the current first "
    "stayed in the warning band (30 A) for 4 s and the 40 A step then needed 1 s of persistence. "
    "The short circuit was tripped in the same 100 ms sample. The CAN communication loss was "
    f"raised by the gateway after {float(next(r for r in prot if r['flag'] == 'COMM_LOSS')['detection_delay_s']):.1f} s "
    "(0.5 s timeout) and cleared as soon as frames returned. In the thermal "
    f"event the dT/dt trend warning came at {float(tr['bms_detect_s']):.0f} s, "
    f"{float(otw['bms_detect_s']) - float(tr['bms_detect_s']):.0f} s before the 50 °C warning and "
    f"{float(otf['bms_detect_s']) - float(tr['bms_detect_s']):.0f} s before the 60 °C fault, "
    "giving the operator an early warning.")
figure("web", FIG / "dashboard_web_fault.jpg",
       "Web dashboard during the over-temperature fault: fault banner, contactor open, active alarms",
       width=14.0)
figure("web2", FIG / "dashboard_web_lowsoc.jpg",
       "Web dashboard at the end of the drive: low SoC, weak cell 3 shown in red, warning alarms",
       width=14.0)
figure("nr", FIG / "dashboard_nodered.jpg",
       "Node-RED dashboard at the end of the scenario: gauges, trend, cell summary and alarm table",
       width=14.0)

heading("11.2 Latency and data integrity", 2)
figure("lat", FIG / "latency_hist.png", "End-to-end latency from CAN frame to MQTT subscriber", width=13.5)
table("latt", "Pipeline performance over the 3000 s scenario", ["Quantity", "Value"], [
    ["CAN frames sent / received by gateway", f"{gw['frames']:,} (56 frames/s)"],
    ["Checksum errors / counter jumps", f"{gw['checksum_errors']} / {gw['counter_jumps']}"],
    ["Communication-loss events detected", f"{gw['comm_loss_events']} (1 injected)"],
    ["Telemetry messages / alarm messages", f"{summary['telemetry_msgs']:,} / {summary['alarm_msgs']}"],
    ["Mean JSON telemetry size", f"{summary['mean_json_bytes']:.0f} bytes"],
    ["Latency median / 95th / 99th percentile", f"{lat['p50']:.1f} / {lat['p95']:.1f} / {lat['p99']:.1f} ms"],
    ["Latency maximum", f"{lat['max']:.0f} ms"],
], [8.5, 7.5])
para(
    f"Half of the telemetry messages reached the subscriber within {lat['p50']:.1f} ms of the CAN "
    f"frame, 95 % within {lat['p95']:.1f} ms and the slowest took {lat['max']:.1f} ms "
    f"(Figure {fig_no['lat']}). Alarms were as fast as telemetry because they are published as soon "
    "as the fault frame arrives. All parts ran on one computer, so on a real link the network "
    "delay (tens to hundreds of milliseconds on cellular) would dominate. "
    "No checksum errors or counter jumps occurred, as expected on a virtual bus.")

heading("11.3 CAN bus load", 2)
cr = [[r["can_id"], r["message"], r["cycle_ms"], r["bits_mean"], r["load_500k_%"]] for r in canmsg]
cr.append(["", "Total", "", "", f"{load500:.2f}"])
table("load", "Frame length (with measured stuff bits) and bus load at 500 kbit/s",
      ["ID", "Message", "Cycle (ms)", "Mean bits/frame", "Load (%)"], cr,
      [2.0, 5.0, 2.5, 3.2, 2.5], size=9.5)
para(
    f"The whole BMS uses only {load500:.2f} % of a 500 kbit/s bus ({load125:.1f} % at 125 kbit/s). "
    "Even with worst-case stuffing (135 bits per frame) it stays below 1.6 %. A shared vehicle "
    "bus is usually kept below about 50 % load, so the same message scheme could report several "
    "hundred cells. A bit-level CRC-15 test (analysis/can_protocol_analysis.py) also confirmed that "
    "CAN detects every error of up to five flipped bits in a frame, while a simple 8-bit checksum "
    "missed up to 6 % of two-bit errors.")

heading("11.4 MQTT quality of service", 2)
figure("qos", FIG / "mqtt_qos.png", "MQTT QoS test: (a) latency, (b) messages delivered after a 3 s outage",
       width=14.5)
table("qost", "MQTT QoS comparison (1000 messages of ~620 bytes; 3 s outage test with 200 messages)",
      ["QoS", "Median latency (ms)", "Packets per message", "Sent during outage", "Lost"], [
    [q[0], f"{float(q[4]):.2f}", q[8], o[2], o[5]] for q, o in zip(qos_lat, qos_out)],
      [1.5, 3.4, 3.4, 3.4, 2.0])
para(
    "QoS 1 and QoS 2 cost less than 1 ms of extra latency on a local broker. During a 3 s "
    "outage of the subscriber, QoS 0 lost every message sent in that time, while QoS 1 and 2 "
    "delivered all of them after reconnection because the broker kept them for the persistent "
    "session. Telemetry is therefore sent with QoS 0 (the next update replaces a lost one), and "
    "alarms with QoS 1 so that no alarm is missed. QoS 2 is not needed, because a duplicated alarm "
    "is harmless.")

heading("11.5 Protocol comparison", 2)
figure("prot", FIG / "protocol_transfer.png",
       "Time needed to transfer one 48-byte BMS snapshot over each link", width=14.0)
table("ptab", "Comparison of communication links for BMS use", [
    "Link", "Data rate", "Range / topology", "Snapshot time", "Typical BMS use"], [
    ["CAN 2.0 / CAN FD", "0.125–1 Mbit/s (FD up to 5 Mbit/s)", "~40 m at 1 Mbit/s; multi-master bus",
     f"{pm['CAN 2.0A 500 kbit/s']:.2f} ms ({pm['CAN FD 0.5/2 Mbit/s']:.2f} ms FD)",
     "EV/vehicle network, BMS ↔ VCU and charger; master ↔ slave boards"],
    ["SMBus", "10–100 kHz", "On one PCB / short cable; master–slave", f"{pm['SMBus 100 kHz']:.2f} ms",
     "Smart batteries (laptops, tools); battery ↔ host charger"],
    ["UART / RS-485", "9.6–115.2 kbit/s", "Point-to-point (RS-485: multi-drop, ~1 km)",
     f"{pm['UART/RS-485 115.2 kbit/s']:.2f} ms", "Debug port, Modbus to inverters, daisy-chained modules"],
    ["Bluetooth LE", "1–2 Mbit/s air rate", "10–50 m; star (phone ↔ BMS)",
     f"{pm['BLE 4.2 1M PHY (air)']:.2f} ms air (+ 7.5–50 ms connection interval)",
     "Phone apps for e-bikes, scooters, home batteries"],
    ["LoRa / LoRaWAN", "0.3–5.5 kbit/s", "2–15 km; star via gateway",
     f"{pm['LoRa SF7 125 kHz']:.0f} ms (SF7) – {pm['LoRa SF12 125 kHz'] / 1000:.1f} s (SF12)",
     "Remote monitoring of stationary storage, telecom and solar sites"],
], [2.6, 2.9, 3.4, 3.4, 4.0], size=8.5)
para(
    f"CAN moves a complete snapshot in {pm['CAN 2.0A 500 kbit/s']:.1f} ms with strong error "
    "detection and priority-based arbitration, which is why it is the standard for the BMS "
    "protection network in vehicles. SMBus and UART are slower but cheap and sufficient for "
    "single-battery products. BLE has a high air rate, but the update rate is set by the "
    "connection interval, and it is a convenience link to a phone rather than a safety link. "
    f"LoRa needs {pm['LoRa SF12 125 kHz'] / 1000:.1f} s of air time per snapshot at SF12; with a "
    "1 % duty-cycle limit this allows one update every few minutes, which is enough for remote "
    "status but not for protection.")

section(12, "OBSERVATION AND DISCUSSION")
rich([("Communication must carry protection, not only data. ", True),
      ("The BMS acted locally (contactor opened in the same 100 ms cycle for the short circuit); the "
       "network only reported the action. Safety functions should never depend on an IoT link.", False)])
rich([("Integrity checks make the data trustworthy. ", True),
      ("The rolling counter, checksum and watchdog let the gateway tell 'battery healthy' apart from "
       "'no data'. Without the watchdog, the dashboard would have kept showing the last values "
       "during the 5 s communication loss.", False)])
rich([("Debounce is a trade-off. ", True),
      ("A 0.5–1 s debounce avoided false alarms from noise and short current peaks, but added the "
       "same delay to real faults. Instantaneous trips are kept for short circuits.", False)])
rich([("Trend alarms give early warning. ", True),
      (f"The dT/dt alarm warned {float(otf['bms_detect_s']) - float(tr['bms_detect_s']):.0f} s before "
       "the over-temperature fault, which is the kind of early thermal-runaway indication that "
       "remote monitoring is most useful for.", False)])
rich([("Bandwidth is not a problem on CAN, but it is on the IoT side. ", True),
      (f"The BMS used {load500:.1f} % of the CAN bus. The JSON telemetry, however, is about "
       f"{summary['mean_json_bytes']:.0f} bytes per second, roughly ten times a binary snapshot, "
       "which adds up to about 1.3 GB per month per pack on a cellular link. A real product would "
       "send binary data or reduce the update rate when nothing changes.", False)])
rich([("Choice of QoS matters. ", True),
      ("QoS 1 for alarms and a retained status topic with a last-will message gave reliable "
       "alarm delivery at almost no cost.", False)])
para("These results come from a simulation on one computer: real bus faults, network delays and "
     "packet loss were not present, so absolute latencies on a real system will be higher.")

section(13, "FUTURE SCOPE")
para(
    "The gateway can be run against a real CAN bus by changing the python-can interface to a USB-CAN "
    "adapter, and the broker can be moved to a cloud server with TLS and user authentication. The "
    "dashboard could store data in a time-series database for long-term SoH trends, and a "
    "two-way channel could let an operator reset latched faults remotely with proper "
    "authorisation. CAN FD and J1939-style multi-packet messages could be used for larger packs, "
    "and the LoRa or BLE links could be tested on low-cost development boards.")

section(14, "CONCLUSION")
para(
    "A complete BMS communication and IoT monitoring chain was built and tested in Python. A pack "
    "model with injected faults fed a BMS ECU that sent eight DBC-defined CAN messages on a "
    "python-can virtual bus. A gateway checked counters, checksums and timeouts and published "
    "telemetry and alarms over MQTT to a web dashboard and a Node-RED dashboard. All seven "
    f"injected events were detected and shown on the dashboards, with a median CAN-to-dashboard "
    f"latency of {lat['p50']:.1f} ms. The BMS used only {load500:.2f} % of a 500 kbit/s CAN bus. "
    "MQTT QoS 1 was needed so that alarms are not lost during network outages. CAN is the right "
    "choice for the on-board BMS network, SMBus and UART for simple single-battery products, and "
    "BLE and LoRa for monitoring only.")

section(15, "REFERENCES")
refs = [
    "D. Andrea, Battery Management Systems for Large Lithium-Ion Battery Packs. Norwood, MA, USA: Artech House, 2010.",
    "R. Xiong and W. Shen, Advanced Battery Management Technologies for Electric Vehicles. Hoboken, NJ, USA: Wiley, 2019.",
    "ISO 11898-1:2015, Road vehicles – Controller area network (CAN) – Part 1: Data link layer and physical signalling.",
    "Robert Bosch GmbH, CAN Specification, Version 2.0, 1991.",
    "R. I. Davis, A. Burns, R. J. Bril and J. J. Lukkien, “Controller Area Network (CAN) schedulability analysis: "
    "Refuted, revisited and revised,” Real-Time Systems, vol. 35, no. 3, pp. 239–272, 2007.",
    "OASIS, MQTT Version 3.1.1, OASIS Standard, 2014.",
    "AUTOSAR, Specification of SW-C End-to-End Communication Protection Library, AUTOSAR Classic Platform.",
    "SBS Implementers Forum, Smart Battery Data Specification, Revision 1.1, 1998.",
    "SBS Implementers Forum, System Management Bus (SMBus) Specification, Version 3.0, 2014.",
    "Bluetooth SIG, Bluetooth Core Specification, Version 4.2, 2014.",
    "Semtech Corporation, AN1200.13: LoRa Modem Designer's Guide, 2013.",
    "P. Koopman and T. Chakravarty, “Cyclic redundancy code (CRC) polynomial selection for embedded networks,” "
    "in Proc. Int. Conf. Dependable Systems and Networks (DSN), 2004, pp. 145–154.",
    "python-can developers, python-can documentation, version 4.6. [Online]. Available: https://python-can.readthedocs.io",
    "OpenJS Foundation, Node-RED documentation. [Online]. Available: https://nodered.org/docs",
    "23RF16 Battery Management Systems – Mini Project Topics, Project 7: BMS Communication and IoT-Based "
    "Battery Monitoring Dashboard (problem statement).",
]
for i, rtext in enumerate(refs, 1):
    p = para(f"[{i}] {rtext}", size=10.5, after=3)

section(16, "ANNEXURE – PYTHON CODE")
para("The listing shows the core of the project: the checksum, the encoding of the fault message "
     "in the BMS ECU, and the gateway's frame handler and watchdog. The full code is in the "
     "project folder (see README.md).")
code('''
import can, cantools, json, time
import paho.mqtt.client as mqtt

db = cantools.database.load_file("can_db/bms.dbc")

def checksum(data, can_id):          # used on BMS_PackStatus and BMS_Faults
    return 0xFF ^ ((sum(data[:7]) + (can_id & 0xFF)) & 0xFF)

# ---------------- BMS ECU: encode and send the fault frame ----------------
def send_faults(bus, warn_flags, fault_flags, counter, fault_count):
    msg = db.get_message_by_name("BMS_Faults")
    data = bytearray(msg.encode(dict(WarningFlags=warn_flags, FaultFlags=fault_flags,
                                     BalancingMask=0, RollingCounter=counter,
                                     FaultCount=fault_count, Checksum=0)))
    data[7] = checksum(data, msg.frame_id)
    bus.send(can.Message(arbitration_id=msg.frame_id, data=data, is_extended_id=False))

# ---------------- Gateway: CAN -> MQTT ----------------
bus = can.Bus(interface="virtual", channel="bms_vcan0")
mq = mqtt.Client(mqtt.CallbackAPIVersion.VERSION2, client_id="bms-can-gateway")
mq.will_set("bms/pack01/status", json.dumps({"online": False}), qos=1, retain=True)
mq.connect("127.0.0.1", 1883); mq.loop_start()

prev_warn = prev_fault = 0
last_status = time.time()
while True:
    frame = bus.recv(timeout=0.01)
    if frame is None:                                  # watchdog when no traffic
        if time.time() - last_status > 0.5:
            mq.publish("bms/pack01/alarms", json.dumps(
                {"code": "COMM_LOSS", "level": "FAULT", "active": True}), qos=1)
        continue
    m = db.get_message_by_frame_id(frame.arbitration_id)
    sig = m.decode(frame.data, decode_choices=False)
    if "Checksum" in sig and checksum(frame.data, frame.arbitration_id) != frame.data[7]:
        continue                                       # corrupted frame -> discard
    if m.name == "BMS_PackStatus":
        last_status = time.time()
        mq.publish("bms/pack01/telemetry", json.dumps(
            {"pack_v": sig["PackVoltage"], "pack_i": sig["PackCurrent"], "soc": sig["SoC"]}))
    elif m.name == "BMS_Faults":
        changed = (sig["WarningFlags"] ^ prev_warn) | (sig["FaultFlags"] ^ prev_fault)
        if changed:                                    # publish an alarm event on change
            mq.publish("bms/pack01/alarms", json.dumps(
                {"warn": sig["WarningFlags"], "fault": sig["FaultFlags"]}), qos=1)
        prev_warn, prev_fault = sig["WarningFlags"], sig["FaultFlags"]
''')

doc.save(OUT)
print("saved", OUT)
